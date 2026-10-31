import csv, sqlite3, threading, calendar, ipaddress
from html import escape
from datetime import datetime, date, timedelta
from pathlib import Path
import tkinter as tk
from tkinter import ttk, messagebox, filedialog

NEW_BASE = Path.home() / "SCAFAWorkforce"
LEGACY_BASE = Path.home() / "K50AttendanceManager"
# Keep using the legacy database when it exists, so upgrades do not silently start
# with an empty database. If a SCAFA Workforce database already exists, prefer it.
if (NEW_BASE / "attendance.db").exists() or not (LEGACY_BASE / "attendance.db").exists():
    BASE = NEW_BASE
else:
    BASE = LEGACY_BASE
BASE.mkdir(parents=True, exist_ok=True)
DB = BASE / "attendance.db"
BG = "#f5f7f4"
NAVY = "#17251d"
BLUE = "#2f6b45"
TEAL = "#477c57"
MUTED = "#68756b"

def init_db():
    with sqlite3.connect(DB) as c:
        c.execute("""CREATE TABLE IF NOT EXISTS punches (
            user_id TEXT NOT NULL, punch_time TEXT NOT NULL, status TEXT DEFAULT '',
            UNIQUE(user_id,punch_time,status))""")
        c.execute("""CREATE TABLE IF NOT EXISTS staff (
            user_id TEXT PRIMARY KEY, name TEXT NOT NULL DEFAULT '')""")
        c.execute("""CREATE TABLE IF NOT EXISTS staff_rules (
            user_id TEXT PRIMARY KEY, start_time TEXT NOT NULL DEFAULT '10:00',
            end_time TEXT NOT NULL DEFAULT '18:00', grace_minutes INTEGER NOT NULL DEFAULT 10,
            half_day_after TEXT NOT NULL DEFAULT '12:00', min_hours REAL NOT NULL DEFAULT 4.0)""")
        c.execute("""CREATE TABLE IF NOT EXISTS app_settings (
            setting TEXT PRIMARY KEY, value TEXT NOT NULL)""")
        c.execute("""CREATE TABLE IF NOT EXISTS holidays (
            holiday_date TEXT PRIMARY KEY, label TEXT NOT NULL DEFAULT '')""")
        c.execute("""CREATE TABLE IF NOT EXISTS staff_day_rules (
            user_id TEXT NOT NULL, weekday INTEGER NOT NULL, start_time TEXT NOT NULL,
            end_time TEXT NOT NULL, grace_minutes INTEGER NOT NULL DEFAULT 15,
            shift_type TEXT NOT NULL DEFAULT 'Full day', PRIMARY KEY(user_id,weekday))""")
        seed = {"weekday_window_start":"09:00","weekday_window_end":"20:00",
                "saturday_window_start":"09:00","saturday_window_end":"17:00"}
        for k,v in seed.items():
            c.execute("INSERT OR IGNORE INTO app_settings(setting,value) VALUES(?,?)",(k,v))

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("SCAFA Workforce")
        self.geometry("1280x840")
        self.minsize(1040, 700)
        self.configure(bg=BG)
        init_db()
        self.ip=tk.StringVar(value="192.168.1.201")
        self.port=tk.StringVar(value="4370")
        self.key=tk.StringVar(value="0")
        self.status=tk.StringVar(value="Ready")
        self.start_date=tk.StringVar(value=date.today().replace(day=1).isoformat())
        self.end_date=tk.StringVar(value=date.today().isoformat())
        self.report_user=tk.StringVar(value="All staff")
        self._style()
        self._ui()
        self.refresh()

    def _date_field(self, parent, variable, width=12):
        wrap=ttk.Frame(parent)
        ttk.Entry(wrap,textvariable=variable,width=width).pack(side="left")
        ttk.Button(wrap,text="▦",width=3,command=lambda:self._open_calendar(variable)).pack(side="left",padx=(3,0))
        return wrap

    def _open_calendar(self, variable):
        try:
            selected=datetime.strptime(variable.get().strip(),"%Y-%m-%d").date()
        except ValueError:
            selected=date.today()
        popup=tk.Toplevel(self); popup.title("Select date"); popup.resizable(False,False); popup.transient(self); popup.grab_set()
        state={"year":selected.year,"month":selected.month}
        head=ttk.Frame(popup,padding=8); head.pack(fill="x")
        grid=ttk.Frame(popup,padding=(8,0,8,8)); grid.pack()
        def draw():
            for child in grid.winfo_children(): child.destroy()
            for child in head.winfo_children(): child.destroy()
            ttk.Button(head,text="‹",width=3,command=lambda:move(-1)).pack(side="left")
            ttk.Label(head,text=f"{calendar.month_name[state['month']]} {state['year']}",font=("Segoe UI",11,"bold")).pack(side="left",expand=True,padx=12)
            ttk.Button(head,text="›",width=3,command=lambda:move(1)).pack(side="right")
            for col,name in enumerate(("Mo","Tu","We","Th","Fr","Sa","Su")):
                ttk.Label(grid,text=name,width=4,anchor="center").grid(row=0,column=col,padx=1,pady=3)
            for row,week in enumerate(calendar.monthcalendar(state["year"],state["month"]),1):
                for col,daynum in enumerate(week):
                    if daynum:
                        ttk.Button(grid,text=str(daynum),width=4,command=lambda d=daynum:choose(d)).grid(row=row,column=col,padx=1,pady=1)
            ttk.Button(popup,text="Today",command=choose_today).pack(pady=(0,8))
        def move(delta):
            month=state["month"]+delta
            if month<1: state["year"]-=1; month=12
            elif month>12: state["year"]+=1; month=1
            state["month"]=month; draw()
        def choose(daynum):
            variable.set(date(state["year"],state["month"],daynum).isoformat()); popup.destroy()
        def choose_today():
            variable.set(date.today().isoformat()); popup.destroy()
        draw()

    def _valid_ip_edit(self, proposed):
        return len(proposed)<=15 and all(ch in "0123456789." for ch in proposed)

    def _style(self):
        s=ttk.Style(self)
        try: s.theme_use("clam")
        except Exception: pass
        s.configure("TFrame", background=BG)
        s.configure("TLabel", background=BG, foreground=NAVY, font=("Segoe UI",10))
        s.configure("Title.TLabel", font=("Segoe UI",22,"bold"), foreground=NAVY)
        s.configure("Sub.TLabel", font=("Segoe UI",10), foreground=MUTED)
        s.configure("Card.TFrame", background="white", relief="flat")
        s.configure("CardTitle.TLabel", background="white", foreground=MUTED, font=("Segoe UI",10))
        s.configure("CardValue.TLabel", background="white", foreground=NAVY, font=("Segoe UI",22,"bold"))
        s.configure("Accent.TButton", font=("Segoe UI",10,"bold"), padding=(12,8), background=BLUE, foreground="white")
        s.map("Accent.TButton", background=[("active","#1d4ed8")])
        s.configure("Treeview", rowheight=34, font=("Segoe UI",9), background="white", fieldbackground="white", borderwidth=0)
        s.configure("Treeview.Heading", font=("Segoe UI",9,"bold"), background="#e5eee7", foreground=NAVY, padding=(8,9))
        s.configure("TNotebook", background=BG, borderwidth=0)
        s.configure("TNotebook.Tab", padding=(20,11), font=("Segoe UI",10,"bold"))
        s.configure("TEntry", padding=6)
        s.configure("TCombobox", padding=6)
        s.configure("TButton", padding=(10,7), font=("Segoe UI",9))

    def _ui(self):
        header=tk.Frame(self,bg=NAVY,height=94); header.pack(fill="x")
        header.pack_propagate(False)
        tk.Label(header,text="SCAFA  |  WORKFORCE",bg="#2f6b45",fg="white",font=("Segoe UI",21,"bold")).pack(anchor="w",padx=24,pady=(15,2))
        tk.Label(header,text="Staff attendance  •  K50 synchronisation  •  Management reporting",bg="#2f6b45",fg="#cbd5e1",font=("Segoe UI",10)).pack(anchor="w",padx=26)
        conn=ttk.Frame(self,padding=(18,14,18,10)); conn.pack(fill="x")
        ttk.Label(conn,text="DEVICE CONNECTION",font=("Segoe UI",9,"bold")).grid(row=0,column=0,columnspan=8,sticky="w",pady=(0,6))
        for col,label,var,width in [(0,"Device IP",self.ip,17),(2,"Port",self.port,7),(4,"Comm key",self.key,7)]:
            ttk.Label(conn,text=label).grid(row=1,column=col,sticky="w",padx=(0,5))
            if label=="Device IP":
                vcmd=(self.register(self._valid_ip_edit),"%P")
                ttk.Entry(conn,textvariable=var,width=18,validate="key",validatecommand=vcmd).grid(row=1,column=col+1,sticky="w",padx=(0,16))
            else:
                ttk.Entry(conn,textvariable=var,width=width).grid(row=1,column=col+1,sticky="w",padx=(0,16))
        ttk.Button(conn,text="⟳  Fetch from K50",style="Accent.TButton",command=self.fetch).grid(row=1,column=6,padx=(0,8))
        ttk.Button(conn,text="＋  Import CSV",command=self.import_csv).grid(row=1,column=7)
        ttk.Button(conn,text="⚙  Custom Rules",command=self.open_rules).grid(row=1,column=8,padx=(8,0))
        ttk.Button(conn,text="Staff Schedules",command=self.open_staff_schedules).grid(row=1,column=9,padx=(8,0))
        ttk.Button(conn,text="Backup Database",command=self.backup_database).grid(row=2,column=6,sticky="w",pady=(8,0))
        ttk.Button(conn,text="Restore Database",command=self.restore_database).grid(row=2,column=7,sticky="w",pady=(8,0))
        filters=ttk.Frame(self,padding=(18,0,18,10)); filters.pack(fill="x")
        ttk.Label(filters,text="From").pack(side="left")
        self._date_field(filters,self.start_date).pack(side="left",padx=(6,14))
        ttk.Label(filters,text="To").pack(side="left")
        self._date_field(filters,self.end_date).pack(side="left",padx=(6,14))
        ttk.Button(filters,text="Apply dates",command=self.refresh).pack(side="left")
        ttk.Label(filters,text="PDF scope").pack(side="left",padx=(22,6))
        self.user_combo=ttk.Combobox(filters,textvariable=self.report_user,width=18,state="readonly",values=["All staff"])
        self.user_combo.pack(side="left")
        ttk.Button(filters,text="Export PDF report",style="Accent.TButton",command=self.export_pdf).pack(side="right")
        ttk.Button(filters,text="Export Excel",command=self.export_excel).pack(side="right",padx=8)
        ttk.Button(filters,text="Export CSV",command=self.export_csv).pack(side="right",padx=8)
        self.tabs=ttk.Notebook(self); self.tabs.pack(fill="both",expand=True,padx=18,pady=(0,8))
        self.dashboard=ttk.Frame(self.tabs,padding=12); self.records=ttk.Frame(self.tabs,padding=12)
        self.tabs.add(self.dashboard,text="  ◈ Dashboard  "); self.tabs.add(self.records,text="  ☷ Punch records  ")
        self._dashboard_ui(); self._records_ui()
        footer=tk.Frame(self,bg="#e8eef8",height=30); footer.pack(fill="x",side="bottom")
        tk.Label(footer,textvariable=self.status,bg="#e8eef8",fg=MUTED,font=("Segoe UI",9),anchor="w").pack(fill="x",padx=18,pady=6)

    def _dashboard_ui(self):
        self.cards=ttk.Frame(self.dashboard); self.cards.pack(fill="x",pady=(0,12))
        self.card_values=[]
        for i,(title,icon) in enumerate([("TOTAL PUNCHES","↗"),("UNIQUE STAFF","♙"),("ACTIVE DAYS","▦"),("AVG. PUNCHES / DAY","◷")]):
            card=ttk.Frame(self.cards,style="Card.TFrame",padding=14); card.grid(row=0,column=i,sticky="nsew",padx=5)
            self.cards.columnconfigure(i,weight=1)
            ttk.Label(card,text=icon+"  "+title,style="CardTitle.TLabel").pack(anchor="w")
            val=ttk.Label(card,text="0",style="CardValue.TLabel"); val.pack(anchor="w",pady=(8,0))
            self.card_values.append(val)
        chartrow=ttk.Frame(self.dashboard); chartrow.pack(fill="both",expand=True)
        self.daily_box=ttk.Frame(chartrow,style="Card.TFrame",padding=12); self.daily_box.pack(side="left",fill="both",expand=True,padx=5)
        self.staff_box=ttk.Frame(chartrow,style="Card.TFrame",padding=12); self.staff_box.pack(side="left",fill="both",expand=True,padx=5)
        ttk.Label(self.daily_box,text="PUNCH ACTIVITY BY DAY",font=("Segoe UI",11,"bold"),background="white").pack(anchor="w")
        ttk.Label(self.staff_box,text="PUNCHES BY STAFF ID",font=("Segoe UI",11,"bold"),background="white").pack(anchor="w")
        self.daily_canvas=tk.Canvas(self.daily_box,bg="white",highlightthickness=0,height=310); self.daily_canvas.pack(fill="both",expand=True,pady=(8,0))
        self.staff_canvas=tk.Canvas(self.staff_box,bg="white",highlightthickness=0,height=310); self.staff_canvas.pack(fill="both",expand=True,pady=(8,0))
        ttk.Label(self.dashboard,text="Charts and statistics reflect imported/fetched punch records in the selected date range. Staff are identified by device ID.",style="Sub.TLabel").pack(anchor="w",pady=(10,0))

    def _records_ui(self):
        cols=("uid","name","time","status")
        self.tree=ttk.Treeview(self.records,columns=cols,show="headings")
        # Keep the data columns compact and predictable instead of spreading
        # every field across the full monitor width.
        specs=[("uid","Device User ID",112),("name","Staff name",230),
               ("time","Attendance date & time",205),("status","Punch status",112)]
        for col,title,width in specs:
            self.tree.heading(col,text=title,anchor="w")
            self.tree.column(col,width=width,minwidth=90,stretch=False,anchor="w")
        self.tree.pack(fill="both",expand=True,padx=2,pady=2)
        self.tree.tag_configure("odd",background="#f4f7f5")
        self.tree.tag_configure("even",background="#ffffff")

    def _rules(self):
        defaults={"default_start":"09:45","default_end":"20:00","default_grace":"15",
                  "half_day_after":"12:00","min_hours":"4.0","workdays":"0,1,2,3,4,5",
                  "weekday_window_start":"09:00","weekday_window_end":"20:00",
                  "saturday_window_start":"09:00","saturday_window_end":"17:00"}
        with sqlite3.connect(DB) as c:
            saved=dict(c.execute("SELECT setting,value FROM app_settings"))
            rules={k:saved.get(k,v) for k,v in defaults.items()}
            staff_rules={str(r[0]):{"start":r[1],"end":r[2],"grace":int(r[3]),"half":r[4],"min_hours":float(r[5])}
                         for r in c.execute("SELECT user_id,start_time,end_time,grace_minutes,half_day_after,min_hours FROM staff_rules")}
            holidays=dict(c.execute("SELECT holiday_date,label FROM holidays"))
        return rules,staff_rules,holidays

    def open_rules(self):
        rules,staff_rules,holidays=self._rules()
        win=tk.Toplevel(self); win.title("Custom Attendance Rules"); win.geometry("900x680"); win.minsize(780,560); win.transient(self)
        outer=ttk.Frame(win,padding=16); outer.pack(fill="both",expand=True)
        ttk.Label(outer,text="CUSTOM ATTENDANCE RULES",font=("Segoe UI",16,"bold")).pack(anchor="w")
        ttk.Label(outer,text="Set defaults, then override schedules for individual staff. Rules apply to attendance reports.",style="Sub.TLabel").pack(anchor="w",pady=(2,12))
        defaults=ttk.LabelFrame(outer,text="Default schedule and thresholds",padding=10); defaults.pack(fill="x")
        vars={}
        fields=[("default_start","Default shift start (HH:MM)"),("default_end","Default shift end (HH:MM)"),("default_grace","Default grace (minutes)"),("half_day_after","Half-day cutoff (HH:MM)"),("min_hours","Minimum hours for full day"),("workdays","Working weekdays (0=Mon, 6=Sun)"),("weekday_window_start","Weekday punch window start"),("weekday_window_end","Weekday punch window end"),("saturday_window_start","Saturday punch window start"),("saturday_window_end","Saturday punch window end")]
        for i,(key,label) in enumerate(fields):
            row=i//3; col=(i%3)*2
            ttk.Label(defaults,text=label).grid(row=row,column=col,sticky="w",padx=(0,5),pady=5)
            vars[key]=tk.StringVar(value=rules[key])
            ttk.Entry(defaults,textvariable=vars[key],width=20).grid(row=row,column=col+1,sticky="ew",padx=(0,12),pady=5)
        for col in range(6): defaults.columnconfigure(col,weight=1)
        staffbox=ttk.LabelFrame(outer,text="Per-staff schedule overrides",padding=8); staffbox.pack(fill="both",expand=True,pady=(12,0))
        ttk.Label(staffbox,text="Select an employee, then set their shift, grace period, half-day cutoff and minimum hours.").pack(anchor="w")
        available=sorted(set(self._staff_names())|set(staff_rules),key=lambda x:(not x.isdigit(),int(x) if x.isdigit() else x.lower()))
        choices=[f"{self._name(uid)} ({uid})" for uid in available]
        selected=tk.StringVar(value=choices[0] if choices else "")
        top=ttk.Frame(staffbox); top.pack(fill="x",pady=8)
        ttk.Label(top,text="Staff").pack(side="left")
        combo=ttk.Combobox(top,textvariable=selected,values=choices,state="readonly",width=34); combo.pack(side="left",padx=8)
        ov={}
        for i,(key,label) in enumerate([("start","Shift start"),("end","Shift end"),("grace","Grace minutes"),("half","Half-day cutoff"),("min_hours","Minimum hours")]):
            ttk.Label(top,text=label).pack(side="left",padx=(6,2))
            ov[key]=tk.StringVar(value="")
            ttk.Entry(top,textvariable=ov[key],width=8).pack(side="left",padx=(0,3))
        def load_staff(*_):
            uid=selected.get().rsplit(" (",1)[-1].rstrip(")") if selected.get() else ""
            r=staff_rules.get(uid,{})
            for k,default_key in [("start","default_start"),("end","default_end"),("grace","default_grace"),("half","half_day_after"),("min_hours","min_hours")]:
                ov[k].set(str(r.get(k,vars[default_key].get())))
        combo.bind("<<ComboboxSelected>>",load_staff)
        def save_override():
            label=selected.get()
            if not label: messagebox.showwarning("Select staff","Fetch staff from K50 first or select a staff member."); return
            uid=label.rsplit(" (",1)[-1].rstrip(")")
            try:
                datetime.strptime(ov["start"].get(),"%H:%M"); datetime.strptime(ov["end"].get(),"%H:%M"); datetime.strptime(ov["half"].get(),"%H:%M")
                grace=int(ov["grace"].get()); hours=float(ov["min_hours"].get())
                if grace<0 or hours<0 or hours>24: raise ValueError()
            except ValueError:
                messagebox.showerror("Invalid rule","Use HH:MM times, non-negative grace minutes, and minimum hours from 0 to 24."); return
            staff_rules[uid]={"start":ov["start"].get(),"end":ov["end"].get(),"grace":grace,"half":ov["half"].get(),"min_hours":hours}
            messagebox.showinfo("Schedule override","Override prepared. Click Save all rules to save changes.")
        ttk.Button(staffbox,text="Save selected staff override",command=save_override).pack(anchor="e",pady=(0,8))
        holidaysbox=ttk.LabelFrame(outer,text="Public holidays / closed dates",padding=8); holidaysbox.pack(fill="x",pady=(12,0))
        hrow=ttk.Frame(holidaysbox); hrow.pack(fill="x")
        hdate=tk.StringVar(); hlabel=tk.StringVar()
        ttk.Label(hrow,text="Date").pack(side="left"); self._date_field(hrow,hdate,width=12).pack(side="left",padx=5)
        ttk.Label(hrow,text="Label").pack(side="left"); ttk.Entry(hrow,textvariable=hlabel,width=24).pack(side="left",padx=5)
        holiday_list=tk.Listbox(holidaysbox,height=4); holiday_list.pack(fill="x",pady=6)
        def refresh_holidays():
            holiday_list.delete(0,"end")
            for d,l in sorted(holidays.items()): holiday_list.insert("end",f"{d} — {l}")
        def add_holiday():
            try: datetime.strptime(hdate.get().strip(),"%Y-%m-%d")
            except ValueError: messagebox.showerror("Invalid date","Use YYYY-MM-DD."); return
            holidays[hdate.get().strip()]=hlabel.get().strip() or "Holiday"; refresh_holidays()
        def remove_holiday():
            sel=holiday_list.curselection()
            if not sel: return
            d=holiday_list.get(sel[0]).split(" — ",1)[0]; holidays.pop(d,None); refresh_holidays()
        hb=ttk.Frame(holidaysbox); hb.pack(fill="x")
        ttk.Button(hb,text="Add holiday",command=add_holiday).pack(side="left")
        ttk.Button(hb,text="Remove selected",command=remove_holiday).pack(side="left",padx=8)
        refresh_holidays(); load_staff()
        def save_all():
            try:
                datetime.strptime(vars["default_start"].get(),"%H:%M"); datetime.strptime(vars["default_end"].get(),"%H:%M"); datetime.strptime(vars["half_day_after"].get(),"%H:%M")
                grace=int(vars["default_grace"].get()); hours=float(vars["min_hours"].get())
                days=[int(x.strip()) for x in vars["workdays"].get().split(",") if x.strip()]
                if grace<0 or hours<0 or hours>24 or not days or any(x<0 or x>6 for x in days): raise ValueError()
            except ValueError:
                messagebox.showerror("Invalid defaults","Check HH:MM times, grace, hours, and weekdays (0–6)."); return
            with sqlite3.connect(DB) as c:
                for k,v in vars.items(): c.execute("INSERT INTO app_settings(setting,value) VALUES(?,?) ON CONFLICT(setting) DO UPDATE SET value=excluded.value",(k,v.get().strip()))
                for uid,r in staff_rules.items():
                    c.execute("""INSERT INTO staff_rules(user_id,start_time,end_time,grace_minutes,half_day_after,min_hours) VALUES(?,?,?,?,?,?)
                                 ON CONFLICT(user_id) DO UPDATE SET start_time=excluded.start_time,end_time=excluded.end_time,grace_minutes=excluded.grace_minutes,half_day_after=excluded.half_day_after,min_hours=excluded.min_hours""",
                              (uid,r["start"],r["end"],r["grace"],r["half"],r["min_hours"]))
                c.execute("DELETE FROM holidays")
                c.executemany("INSERT INTO holidays(holiday_date,label) VALUES(?,?)",sorted(holidays.items()))
            messagebox.showinfo("Rules saved","Custom attendance rules saved on this computer.")
            win.destroy()
        buttons=ttk.Frame(outer); buttons.pack(fill="x",pady=(12,0))
        ttk.Button(buttons,text="Save all rules",style="Accent.TButton",command=save_all).pack(side="right")
        ttk.Button(buttons,text="Close",command=win.destroy).pack(side="right",padx=8)

    def open_staff_schedules(self):
        """Set individual schedules for each employee and weekday."""
        with sqlite3.connect(DB) as c:
            staff = dict(c.execute("SELECT user_id,name FROM staff"))
            existing = {(str(uid),int(wd)):(st,en,int(gr),typ) for uid,wd,st,en,gr,typ in
                c.execute("SELECT user_id,weekday,start_time,end_time,grace_minutes,shift_type FROM staff_day_rules")}
        staff.update({uid:self._name(uid) for uid in self._staff_names()})
        if not staff:
            messagebox.showinfo("No employees","Fetch employee names from the K50 first.")
            return
        win=tk.Toplevel(self); win.title("SCAFA Workforce — Staff Schedules"); win.geometry("850x560"); win.transient(self)
        outer=ttk.Frame(win,padding=16); outer.pack(fill="both",expand=True)
        ttk.Label(outer,text="INDIVIDUAL STAFF SCHEDULES",font=("Segoe UI",16,"bold")).pack(anchor="w")
        ttk.Label(outer,text="Configure each employee's shift and grace period separately for every weekday.",style="Sub.TLabel").pack(anchor="w",pady=(3,12))
        choices=[f"{staff[uid]} ({uid})" for uid in sorted(staff,key=lambda x:staff[x].lower())]
        selected=tk.StringVar(value=choices[0]); top=ttk.Frame(outer); top.pack(fill="x")
        ttk.Label(top,text="Employee").pack(side="left")
        combo=ttk.Combobox(top,textvariable=selected,values=choices,state="readonly",width=40); combo.pack(side="left",padx=8)
        grid=ttk.Frame(outer); grid.pack(fill="both",expand=True,pady=12)
        for j,title in enumerate(("Day","Start (HH:MM)","End (HH:MM)","Grace minutes","Shift type")):
            ttk.Label(grid,text=title,font=("Segoe UI",9,"bold")).grid(row=0,column=j,sticky="w",padx=5,pady=5)
        days=("Monday","Tuesday","Wednesday","Thursday","Friday","Saturday","Sunday")
        cells={}
        for wd,name in enumerate(days):
            ttk.Label(grid,text=name).grid(row=wd+1,column=0,sticky="w",padx=5,pady=4)
            vals=existing.get((selected.get().rsplit(" (",1)[-1].rstrip(")"),wd),("09:45","20:00" if wd<5 else "17:00",15,"Off" if wd==6 else "Full day"))
            sv,ev,gv,tv=tk.StringVar(value=vals[0]),tk.StringVar(value=vals[1]),tk.StringVar(value=str(vals[2])),tk.StringVar(value=vals[3])
            cells[wd]=(sv,ev,gv,tv)
            ttk.Entry(grid,textvariable=sv,width=14).grid(row=wd+1,column=1,padx=5,pady=3)
            ttk.Entry(grid,textvariable=ev,width=14).grid(row=wd+1,column=2,padx=5,pady=3)
            ttk.Entry(grid,textvariable=gv,width=14).grid(row=wd+1,column=3,padx=5,pady=3)
            ttk.Combobox(grid,textvariable=tv,values=("Full day","Half day","Custom","Off"),state="readonly",width=13).grid(row=wd+1,column=4,padx=5,pady=3)
        def load_employee(*_):
            uid=selected.get().rsplit(" (",1)[-1].rstrip(")")
            for wd in range(7):
                vals=existing.get((uid,wd),("09:45","20:00" if wd<5 else "17:00",15,"Off" if wd==6 else "Full day"))
                for var,val in zip(cells[wd],vals): var.set(str(val))
        combo.bind("<<ComboboxSelected>>",load_employee)
        def save():
            uid=selected.get().rsplit(" (",1)[-1].rstrip(")")
            try:
                rows=[]
                for wd,vs in cells.items():
                    st,en,gr,typ=[v.get().strip() for v in vs]
                    datetime.strptime(st,"%H:%M"); datetime.strptime(en,"%H:%M")
                    if int(gr)<0 or typ not in ("Full day","Half day","Custom","Off"): raise ValueError()
                    rows.append((uid,wd,st,en,int(gr),typ))
            except ValueError:
                messagebox.showerror("Invalid schedule","Use HH:MM times, a non-negative grace period, and a valid shift type."); return
            with sqlite3.connect(DB) as c:
                c.executemany("""INSERT INTO staff_day_rules(user_id,weekday,start_time,end_time,grace_minutes,shift_type)
                    VALUES(?,?,?,?,?,?) ON CONFLICT(user_id,weekday) DO UPDATE SET start_time=excluded.start_time,
                    end_time=excluded.end_time,grace_minutes=excluded.grace_minutes,shift_type=excluded.shift_type""",rows)
            for uid0,wd,st,en,gr,typ in rows: existing[(uid0,wd)]=(st,en,gr,typ)
            messagebox.showinfo("Saved","Individual schedules saved.")
        ttk.Label(outer,text="Default windows: Mon–Fri 9:00 AM–8:00 PM; Saturday 9:00 AM–5:00 PM. Change these in Custom Rules.",style="Sub.TLabel").pack(anchor="w")
        ttk.Button(outer,text="Save employee schedule",style="Accent.TButton",command=save).pack(anchor="e")

    def _attendance_day_status(self,uid,day,stamps):
        rules,overrides,holidays=self._rules()
        if day in holidays: return "Holiday", "", "", 0.0
        try: workdays={int(x.strip()) for x in rules["workdays"].split(",") if x.strip()}
        except ValueError: workdays={0,1,2,3,4}
        d=datetime.strptime(day,"%Y-%m-%d").date()
        r=overrides.get(str(uid),{})
        with sqlite3.connect(DB) as c:
            day_rule=c.execute("SELECT start_time,end_time,grace_minutes,shift_type FROM staff_day_rules WHERE user_id=? AND weekday=?",(str(uid),d.weekday())).fetchone()
        # An explicit employee/day schedule takes precedence over the global workweek.
        if day_rule:
            start,end,grace,shift_type=day_rule[0],day_rule[1],int(day_rule[2]),day_rule[3]
            if shift_type=="Off": return "Day off", "", "", 0.0
        else:
            if d.weekday() not in workdays: return "Day off", "", "", 0.0
            start=r.get("start",rules["default_start"]); end=r.get("end",rules["default_end"])
            grace=int(r.get("grace",rules["default_grace"])); shift_type="Full day"
        half=r.get("half",rules["half_day_after"])
        min_hours=float(r.get("min_hours",rules["min_hours"]))
        stamps=sorted(stamps)
        if not stamps: return "Absent", "", "", 0.0
        first=datetime.strptime(stamps[0],"%Y-%m-%d %H:%M:%S")
        last=datetime.strptime(stamps[-1],"%Y-%m-%d %H:%M:%S")
        scheduled=datetime.strptime(day+" "+start,"%Y-%m-%d %H:%M")
        late=(first-scheduled).total_seconds()/60
        worked=max(0,(last-first).total_seconds()/3600)
        if len(stamps)<2:
            status="Incomplete punch"
            if late>grace: status+=" — Late"
            return status,first.strftime("%H:%M"),last.strftime("%H:%M"),worked
        if shift_type=="Half day":
            status="Half-day"
        elif shift_type=="Custom":
            status="Present" if worked>=min_hours else "Half-day"
        elif first.strftime("%H:%M")>=half or worked<min_hours:
            status="Half-day"
        else:
            status="Present"
        if late>grace:
            status = "Late" if status=="Present" else status+" — Late"
        if end and last.strftime("%H:%M") < end:
            status += " (early departure)"
        return status,first.strftime("%H:%M"),last.strftime("%H:%M"),worked

    def _staff_names(self):
        if not hasattr(self,"_names_cache") or self._names_cache is None:
            with sqlite3.connect(DB) as c:
                self._names_cache={str(uid): str(name or "").strip() for uid,name in c.execute("SELECT user_id,name FROM staff")}
        return self._names_cache

    def _name(self, uid):
        name=self._staff_names().get(str(uid), "")
        return name if name else f"User {uid}"

    def _store_device_users(self, users):
        saved=0
        with sqlite3.connect(DB) as c:
            for user in users:
                uid=str(getattr(user,"user_id",None) or getattr(user,"uid","")).strip()
                name=str(getattr(user,"name","") or "").strip()
                if not uid or not name:
                    continue
                c.execute("""INSERT INTO staff(user_id,name) VALUES(?,?)
                             ON CONFLICT(user_id) DO UPDATE SET name=excluded.name""",(uid,name))
                saved+=1
        self._names_cache=None
        return saved

    def _date_filter(self):
        try:
            start=datetime.strptime(self.start_date.get().strip(),"%Y-%m-%d").date()
            end=datetime.strptime(self.end_date.get().strip(),"%Y-%m-%d").date()
            if end < start: raise ValueError("End date must not be before start date.")
            return start.isoformat()+" 00:00:00",end.isoformat()+" 23:59:59"
        except ValueError as e:
            messagebox.showerror("Check date range",str(e) or "Use YYYY-MM-DD date format.")
            return None

    def _rows(self, filtered=True):
        with sqlite3.connect(DB) as c:
            if filtered:
                bounds=self._date_filter()
                if not bounds: return []
                return c.execute("SELECT user_id,punch_time,status FROM punches WHERE punch_time BETWEEN ? AND ? ORDER BY punch_time",(bounds[0],bounds[1])).fetchall()
            return c.execute("SELECT user_id,punch_time,status FROM punches ORDER BY punch_time").fetchall()

    def refresh(self):
        rows=self._rows()
        if rows is None: return
        for i in self.tree.get_children(): self.tree.delete(i)
        for idx,row in enumerate(reversed(rows[-5000:])):
            self.tree.insert("","end",values=(row[0],self._name(row[0]),row[1],row[2]),tags=("odd",) if idx%2 else ())
        with sqlite3.connect(DB) as c:
            registered={str(r[0]) for r in c.execute("SELECT user_id FROM staff")}
        users=sorted(registered | set(str(r[0]) for r in rows),key=lambda x:(not x.isdigit(),int(x) if x.isdigit() else x.lower()))
        choices=["All staff"]+[f"{self._name(uid)} ({uid})" for uid in users]
        self.user_combo["values"]=choices
        if self.report_user.get() not in choices: self.report_user.set("All staff")
        days=len(set(r[1][:10] for r in rows))
        staff=len(users)
        avg=(len(rows)/days) if days else 0
        for label,value in zip(self.card_values,[f"{len(rows):,}",f"{staff:,}",f"{days:,}",f"{avg:.1f}"]): label.configure(text=value)
        self._draw_charts(rows)
        self.status.set(f"{len(rows):,} punches • {staff} staff IDs • {self.start_date.get()} to {self.end_date.get()}")

    def _draw_charts(self,rows):
        daily={}
        staff={}
        for uid,stamp,state in rows:
            daily[stamp[:10]]=daily.get(stamp[:10],0)+1
            staff[str(uid)]=staff.get(str(uid),0)+1
        self._bar_chart(self.daily_canvas,sorted(daily.items())[-14:], "#2f6b45", date_labels=True)
        named_staff={}
        for uid,count in staff.items():
            label=self._name(uid)
            named_staff[label]=named_staff.get(label,0)+count
        self._bar_chart(self.staff_canvas,sorted(named_staff.items(),key=lambda x:x[1],reverse=True)[:10], "#477c57", date_labels=False)

    def _bar_chart(self,canvas,items,color,date_labels=False):
        canvas.delete("all")
        w=max(canvas.winfo_width(),360); h=max(canvas.winfo_height(),260)
        if not items:
            canvas.create_text(w/2,h/2,text="No records in this date range",fill=MUTED,font=("Segoe UI",11))
            return
        maxv=max(v for _,v in items) or 1
        if not date_labels:
            # Horizontal bars give each staff member a dedicated label row.
            # This prevents long employee names from colliding on the x-axis.
            left,right,top,bottom=118,34,18,22
            cw=max(80,w-left-right); ch=max(80,h-top-bottom)
            row_h=ch/max(1,len(items)); bar_h=min(20,row_h*.56)
            for n in range(4):
                x=left+cw*n/3
                canvas.create_line(x,top,x,top+ch,fill="#e2e8f0")
                canvas.create_text(x,top-7,text=str(round(maxv*n/3)),anchor="s",
                                   fill=MUTED,font=("Segoe UI",8))
            for i,(label,val) in enumerate(items):
                cy=top+row_h*(i+.5)
                shown=label if len(label)<=17 else label[:16]+"…"
                canvas.create_text(left-10,cy,text=shown,anchor="e",
                                   fill=MUTED,font=("Segoe UI",8))
                bw=cw*val/maxv
                canvas.create_rectangle(left,cy-bar_h/2,left+bw,cy+bar_h/2,
                                        fill=color,outline="")
                canvas.create_text(min(left+bw+7,w-right),cy,text=str(val),
                                   anchor="w",fill=NAVY,font=("Segoe UI",8,"bold"))
            return
        left,right,top,bottom=42,16,18,48
        cw=w-left-right; ch=h-top-bottom
        for n in range(4):
            y=top+ch-(ch*n/3)
            canvas.create_line(left,y,w-right,y,fill="#e2e8f0")
            canvas.create_text(left-8,y,text=str(round(maxv*n/3)),anchor="e",
                               fill=MUTED,font=("Segoe UI",8))
        gap=cw/len(items); bw=max(8,min(34,gap*.58))
        for i,(label,val) in enumerate(items):
            x=left+gap*i+(gap-bw)/2
            bh=ch*val/maxv
            canvas.create_rectangle(x,top+ch-bh,x+bw,top+ch,fill=color,outline="")
            short=label[5:] if len(label)>5 else label
            canvas.create_text(x+bw/2,top+ch+9,text=short,anchor="n",
                               fill=MUTED,font=("Segoe UI",8))
            canvas.create_text(x+bw/2,top+ch-bh-9,text=str(val),
                               fill=NAVY,font=("Segoe UI",8,"bold"))

    def fetch(self):
        ip=self.ip.get().strip()
        try:
            parsed_ip=ipaddress.ip_address(ip)
            if parsed_ip.version != 4: raise ValueError()
        except ValueError:
            messagebox.showerror("Invalid device IP","Enter a valid IPv4 address, for example 192.168.1.201.")
            return
        try: port=int(self.port.get()); key=int(self.key.get() or "0")
        except ValueError:
            messagebox.showerror("Invalid settings","Port and communication key must be numbers."); return
        self.status.set("Connecting to device…")
        threading.Thread(target=self._fetch_worker,args=(ip,port,key),daemon=True).start()

    def _fetch_worker(self,ip,port,key):
        conn=None
        try:
            from zk import ZK
            conn=ZK(ip,port=port,timeout=12,password=key,force_udp=False, ommit_ping=False).connect()
            conn.disable_device()
            device_users=conn.get_users()
            saved_users=self._store_device_users(device_users)
            rows=conn.get_attendance()
            inserted=0
            with sqlite3.connect(DB) as c:
                for row in rows:
                    stamp=row.timestamp.strftime("%Y-%m-%d %H:%M:%S")
                    cur=c.execute("INSERT OR IGNORE INTO punches(user_id,punch_time,status) VALUES(?,?,?)",(str(row.user_id),stamp,str(getattr(row,"status",""))))
                    inserted+=cur.rowcount
            self.after(0,lambda:self._fetch_done(len(rows),inserted,saved_users))
        except Exception as e:
            msg=str(e); self.after(0,lambda m=msg:self._fetch_error(m))
        finally:
            if conn:
                try: conn.enable_device()
                except Exception: pass
                try: conn.disconnect()
                except Exception: pass

    def _fetch_done(self,total,inserted,saved_users=0):
        self.refresh()
        self.status.set(f"Sync complete • {saved_users} staff names synced • {total} punches read • {inserted} new")
        messagebox.showinfo("Sync complete",f"Staff names fetched: {saved_users}\nAttendance records read: {total}\nNew attendance records saved: {inserted}")

    def _fetch_error(self,msg):
        self.status.set("Device connection failed.")        messagebox.showerror("Could not fetch attendance",f"{msg}\n\nCheck the K50 IP, LAN connection, port, and communication key.")

    def import_csv(self):
        path=filedialog.askopenfilename(title="Import attendance CSV",filetypes=[("CSV files","*.csv"),("All files","*.*")])
        if not path: return
        count=0
        try:
            with open(path,"r",newline="",encoding="utf-8-sig") as f:
                sample=f.read(4096); f.seek(0)
                try: dialect=csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
                except csv.Error: dialect=csv.excel
                reader=csv.DictReader(f,dialect=dialect)
                fields={x.strip().lower():x for x in (reader.fieldnames or [])}
                uidcol=next((fields[x] for x in ("user id","userid","pin","employee id","device user id") if x in fields),None)
                timecol=next((fields[x] for x in ("date/time","datetime","timestamp","punch time","date time","time") if x in fields),None)
                statecol=next((fields[x] for x in ("status","state","punch status") if x in fields),None)
                if not uidcol or not timecol: raise ValueError("CSV needs User ID/PIN and Date/Time/Timestamp columns.")
                with sqlite3.connect(DB) as c:
                    for row in reader:
                        uid=(row.get(uidcol) or "").strip(); stamp=(row.get(timecol) or "").strip(); state=(row.get(statecol) or "").strip() if statecol else ""
                        if uid and stamp:
                            try: stamp=datetime.fromisoformat(stamp).strftime("%Y-%m-%d %H:%M:%S")
                            except ValueError: pass
                            count+=c.execute("INSERT OR IGNORE INTO punches(user_id,punch_time,status) VALUES(?,?,?)",(uid,stamp,state)).rowcount
            self.refresh()
            messagebox.showinfo("Import complete",f"Imported {count} new records.")
        except Exception as e: messagebox.showerror("Import failed",str(e))

    def backup_database(self):
        """Export a consistent SQLite backup without copying a live database file."""
        path=filedialog.asksaveasfilename(
            title="Back up SCAFA Workforce database",
            defaultextension=".db",
            filetypes=[("SQLite database","*.db"),("All files","*.*")],
            initialfile=f"SCAFA_Workforce_Backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db")
        if not path:
            return
        try:
            with sqlite3.connect(DB) as source, sqlite3.connect(path) as target:
                source.backup(target)
            messagebox.showinfo("Backup complete",f"Database backup saved to:\n{path}")
            self.status.set(f"Database backup saved • {path}")
        except Exception as e:
            messagebox.showerror("Backup failed",f"Could not back up the database:\n{e}")

    def restore_database(self):
        """Validate a backup, preserve the current database, then restore safely."""
        path=filedialog.askopenfilename(
            title="Select SCAFA Workforce database backup",
            filetypes=[("SQLite database","*.db *.sqlite *.sqlite3"),("All files","*.*")])
        if not path:
            return
        try:
            if Path(path).resolve()==DB.resolve():
                messagebox.showerror("Invalid backup","Choose a backup file, not the database currently in use.")
                return
            with sqlite3.connect(f"file:{Path(path).resolve().as_posix()}?mode=ro",uri=True) as check:
                integrity=check.execute("PRAGMA integrity_check").fetchone()
                tables={r[0] for r in check.execute("SELECT name FROM sqlite_master WHERE type='table'")}
                if not integrity or integrity[0]!="ok":
                    raise ValueError("The selected file failed SQLite's integrity check.")
                if not {"punches","staff"}.issubset(tables):
                    raise ValueError("This file does not look like a SCAFA Workforce attendance database.")
            if not messagebox.askyesno(
                "Confirm database restore",
                "Restoring replaces the current attendance data and settings. "
                "A safety backup of the current database will be created first. Continue?"):
                return
            safety=BASE / f"Pre_Restore_Backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
            with sqlite3.connect(DB) as current, sqlite3.connect(safety) as backup:
                current.backup(backup)
            with sqlite3.connect(f"file:{Path(path).resolve().as_posix()}?mode=ro",uri=True) as source, sqlite3.connect(DB) as target:
                source.backup(target)
            self._names_cache=None
            init_db()
            self.refresh()
            messagebox.showinfo("Restore complete",f"Database restored successfully.\nSafety backup: {safety}")
            self.status.set("Database restored • current database safety backup retained")
        except Exception as e:
            messagebox.showerror("Restore failed",f"Could not restore this database:\n{e}")

    def export_csv(self):
        rows=self._rows()
        if rows is None: return
        path=filedialog.asksaveasfilename(title="Save punch report",defaultextension=".csv",filetypes=[("CSV report","*.csv")],initialfile=f"SCAFA_Workforce_Punches_{date.today().isoformat()}.csv")
        if not path: return
        with open(path,"w",newline="",encoding="utf-8-sig") as f:
            w=csv.writer(f); w.writerow(["Device User ID","Staff Name","Attendance date & time","Punch status"])
            w.writerows((uid,self._name(uid),stamp,state) for uid,stamp,state in rows)
        messagebox.showinfo("CSV exported",f"Exported {len(rows)} records.")

    def _report_scope_rows(self):
        rows=self._rows()
        if rows is None: return None,None
        label=self.report_user.get()
        scope="All staff" if label=="All staff" else label.rsplit(" (",1)[-1].rstrip(")")
        if scope!="All staff": rows=[r for r in rows if str(r[0])==scope]
        with sqlite3.connect(DB) as c:
            registered={str(r[0]) for r in c.execute("SELECT user_id FROM staff")}
        if scope=="All staff" and not (rows or registered):
            messagebox.showwarning("No staff","No attendance punches or registered staff were found. Fetch staff from the K50 first."); return None,None
        if scope!="All staff" and scope not in (registered | {str(r[0]) for r in rows}):
            messagebox.showwarning("Unknown staff","This staff ID is not registered and has no punches in the selected period."); return None,None
        return rows,scope

    def export_excel(self):
        rows,scope=self._report_scope_rows()
        if rows is None: return
        suffix="All_Staff" if scope=="All staff" else f"Staff_{scope}"
        path=filedialog.asksaveasfilename(title="Export editable Excel attendance report",defaultextension=".xlsx",
            filetypes=[("Excel workbook","*.xlsx")],initialfile=f"SCAFA_Workforce_Attendance_{suffix}_{date.today().isoformat()}.xlsx")
        if not path: return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.table import Table as ExcelTable, TableStyleInfo
            wb=Workbook()
            overview=wb.active; overview.title="Overview"
            summary=wb.create_sheet("Attendance Summary")
            raw=wb.create_sheet("Raw Punches")
            summary.append(["SCAFA WORKFORCE | ATTENDANCE REPORT"])
            summary.merge_cells("A1:H1")
            summary["A1"].font=Font(name="Aptos Display",size=18,bold=True,color="FFFFFF")
            summary["A1"].fill=PatternFill("solid",fgColor="173D2B")
            summary["A1"].alignment=Alignment(vertical="center")
            summary.row_dimensions[1].height=34
            summary.append(["Report scope",scope,"From",self.start_date.get(),"To",self.end_date.get(),"Generated",datetime.now().strftime("%d %b %Y %H:%M")])
            summary.append([])
            summary.append(["Device User ID","Staff Name","Date","Punches","First Punch","Last Punch","Hours","Attendance Status"])
            raw.append(["Device User ID","Staff Name","Attendance Date & Time","Punch Status"])
            for uid,stamp,state in rows: raw.append([str(uid),self._name(uid),stamp,state or ""])
            by_user_day={}
            for uid,stamp,state in rows: by_user_day.setdefault((str(uid),stamp[:10]),[]).append(stamp)
            with sqlite3.connect(DB) as c:
                known=sorted({str(r[0]) for r in c.execute("SELECT user_id FROM staff")} | {str(r[0]) for r in rows},
                    key=lambda x:(not x.isdigit(),int(x) if x.isdigit() else x.lower()))
            first=datetime.strptime(self.start_date.get(),"%Y-%m-%d").date()
            last=datetime.strptime(self.end_date.get(),"%Y-%m-%d").date()
            if last<first: raise ValueError("End date must not precede start date.")
            users=known if scope=="All staff" else [str(scope)]
            for uid in users:
                cursor=first
                while cursor<=last:
                    day=cursor.isoformat(); stamps=sorted(by_user_day.get((uid,day),[]))
                    status,first_p,last_p,hours=self._attendance_day_status(uid,day,stamps)
                    summary.append([uid,self._name(uid),day,len(stamps),first_p or "",last_p or "",round(hours,2),status])
                    cursor+=timedelta(days=1)
            overview.merge_cells("A1:F1")
            overview["A1"]="SCAFA WORKFORCE | ATTENDANCE OVERVIEW"
            overview["A1"].font=Font(name="Aptos Display",size=18,bold=True,color="FFFFFF")
            overview["A1"].fill=PatternFill("solid",fgColor="173D2B")
            overview["A1"].alignment=Alignment(vertical="center")
            overview.row_dimensions[1].height=36
            overview.append([])
            overview.append(["Report scope",scope,"From",self.start_date.get(),"To",self.end_date.get()])
            overview.append(["Generated",datetime.now().strftime("%d %b %Y %H:%M")])
            overview.append([])
            overview.append(["KEY METRICS","Value",None,None,"ATTENDANCE STATUS","Days"])
            endrow=summary.max_row
            for label,formula in [
                ("Total punches",f"=SUM('Attendance Summary'!D5:D{endrow})"),
                ("Staff included",f'=SUMPRODUCT((\'Attendance Summary\'!A5:A{endrow}<>"")/COUNTIF(\'Attendance Summary\'!A5:A{endrow},\'Attendance Summary\'!A5:A{endrow}))'),
                ("Active attendance days",f'=COUNTIF(\'Attendance Summary\'!D5:D{endrow},">0")'),
                ("Present / late days",f'=COUNTIF(\'Attendance Summary\'!H5:H{endrow},"Present")+COUNTIF(\'Attendance Summary\'!H5:H{endrow},"Late")')]:
                overview.append([label,formula])
            status_start=7
            for offset,status_name in enumerate(["Present","Late","Half Day","Absent","Day off","Holiday","Incomplete punch"]):
                rr=status_start+offset
                overview.cell(rr,5,status_name)
                overview.cell(rr,6,f'=COUNTIF(\'Attendance Summary\'!H5:H{endrow},E{rr})')
            staff_start=14
            overview.cell(staff_start,1,"STAFF PUNCH TOTALS")
            for j,h in enumerate(["Device User ID","Staff Name","Punches","Active Days"],1):
                overview.cell(staff_start+1,j,h)
            for idx,uid in enumerate(users,staff_start+2):
                overview.cell(idx,1,uid)
                overview.cell(idx,2,self._name(uid))
                overview.cell(idx,3,f'=SUMIF(\'Attendance Summary\'!A$5:A${endrow},A{idx},\'Attendance Summary\'!D$5:D${endrow})')
                overview.cell(idx,4,f'=COUNTIFS(\'Attendance Summary\'!A$5:A${endrow},A{idx},\'Attendance Summary\'!D$5:D${endrow},">0")')
            overview.sheet_view.showGridLines=False
            for col,width in {"A":24,"B":25,"C":16,"D":16,"E":23,"F":14}.items():
                overview.column_dimensions[col].width=width
            for row in overview.iter_rows():
                for cell in row: cell.alignment=Alignment(vertical="center",horizontal="left")
            for rowno in (6,15):
                for cell in overview[rowno]:
                    cell.font=Font(bold=True,color="FFFFFF")
                    cell.fill=PatternFill("solid",fgColor="24563B")
                    cell.alignment=Alignment(horizontal="center",vertical="center")
            overview.freeze_panes="A7"
            overview.sheet_properties.pageSetUpPr.fitToPage=True
            overview.page_setup.orientation="portrait"
            overview.page_setup.paperSize=overview.PAPERSIZE_A4
            overview.page_setup.fitToWidth=1
            overview.page_setup.fitToHeight=1
            for ws,widths,header_row in [(summary,[16,28,14,11,21,21,12,28],4),(raw,[18,30,24,20],1)]:
                ws.freeze_panes="A5" if ws is summary else "A2"
                ws.sheet_view.showGridLines=False
                for i,width in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=width
                for cell in ws[header_row]:
                    cell.font=Font(bold=True,color="FFFFFF")
                    cell.fill=PatternFill("solid",fgColor="24563B")
                    cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
                ws.row_dimensions[header_row].height=30
                ws.auto_filter.ref=ws.dimensions
                for row in ws.iter_rows():
                    for cell in row:
                        cell.alignment=Alignment(vertical="center",horizontal="left",
                                                  wrap_text=(cell.row==header_row))
                for rr in range(header_row+1,ws.max_row+1):
                    ws.row_dimensions[rr].height=21
                if ws is summary:
                    for rr in range(5,ws.max_row+1):
                        ws.cell(rr,7).number_format="0.00"
                else:
                    for rr in range(2,ws.max_row+1):
                        ws.cell(rr,3).number_format="yyyy-mm-dd hh:mm:ss"
            if summary.max_row>=5:
                tab=ExcelTable(displayName="AttendanceDetails",ref=f"A4:H{summary.max_row}")
                tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium4",showRowStripes=True,showColumnStripes=False)
                summary.add_table(tab)
            if raw.max_row>=2:
                tab=ExcelTable(displayName="RawPunchRecords",ref=f"A1:D{raw.max_row}")
                tab.tableStyleInfo=TableStyleInfo(name="TableStyleMedium4",showRowStripes=True,showColumnStripes=False)
                raw.add_table(tab)
            from openpyxl.formatting.rule import FormulaRule
            for label,fill in [("Present","E2F0E7"),("Late","FFF2CC"),("Half Day","FCE4D6"),
                               ("Absent","F4CCCC"),("Day off","E7E6E6"),("Holiday","DDEBF7"),
                               ("Incomplete punch","FCE4D6")]:
                summary.conditional_formatting.add(f"H5:H{summary.max_row}",
                    FormulaRule(formula=[f'$H5="{label}"'],fill=PatternFill("solid",fgColor=fill)))
            summary.sheet_properties.pageSetUpPr.fitToPage=True
            summary.page_setup.orientation="landscape"; summary.page_setup.paperSize=summary.PAPERSIZE_A4
            summary.page_setup.fitToWidth=1; summary.page_setup.fitToHeight=0
            raw.sheet_properties.pageSetUpPr.fitToPage=True
            raw.page_setup.orientation="landscape"; raw.page_setup.paperSize=raw.PAPERSIZE_A4
            raw.page_setup.fitToWidth=1; raw.page_setup.fitToHeight=0
            wb.active=0
            wb.save(path)
            messagebox.showinfo("Excel exported",f"Editable Excel workbook saved to:\n{path}\n\nIncludes daily attendance summary and raw punch records.")
        except Exception as e:
            messagebox.showerror("Excel export failed",f"{e}\n\nPlease reinstall/update the app build.")

    def export_pdf(self):
        rows,scope=self._report_scope_rows()
        if rows is None: return
        suffix="All_Staff" if scope=="All staff" else "Staff_"+scope
        path=filedialog.asksaveasfilename(title="Save PDF attendance report",defaultextension=".pdf",filetypes=[("PDF report","*.pdf")],initialfile=f"SCAFA_Workforce_Attendance_{suffix}_{date.today().isoformat()}.pdf")
        if not path: return
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.enums import TA_LEFT
            from reportlab.lib.units import mm
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
            from reportlab.graphics.shapes import Drawing
            from reportlab.graphics.charts.barcharts import VerticalBarChart
            self._make_pdf(path,rows,scope,colors,A4,getSampleStyleSheet,ParagraphStyle,TA_LEFT,mm,SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak,Drawing,VerticalBarChart)            messagebox.showinfo("PDF report created",f"Saved report to:\n{path}")
        except Exception as e:
            messagebox.showerror("PDF export failed",f"{e}\n\nIf this persists, reinstall/update the app build.")

    def _make_pdf(self,path,rows,scope,colors,A4,getStyles,ParagraphStyle,TA_LEFT,mm,Doc,Paragraph,Spacer,Table,TableStyle,PageBreak,Drawing,VerticalBarChart):
        styles=getStyles()
        styles.add(ParagraphStyle(name="ReportTitle",parent=styles["Title"],textColor=colors.HexColor("#173D2B"),fontName="Helvetica-Bold",fontSize=18,leading=22,alignment=TA_LEFT,spaceAfter=4))
        styles.add(ParagraphStyle(name="EmployeeTitle",parent=styles["Title"],textColor=colors.HexColor("#173D2B"),fontName="Helvetica-Bold",fontSize=16,leading=20,alignment=TA_LEFT,spaceAfter=4,wordWrap="CJK"))
        styles.add(ParagraphStyle(name="SmallNote",parent=styles["BodyText"],fontSize=7,textColor=colors.HexColor("#64748b"),leading=9))
        styles.add(ParagraphStyle(name="Cell",parent=styles["BodyText"],fontSize=6.7,leading=8,wordWrap="CJK"))
        styles.add(ParagraphStyle(name="CellHead",parent=styles["BodyText"],fontName="Helvetica-Bold",fontSize=6.5,leading=7.5,textColor=colors.white,wordWrap="CJK"))
        styles.add(ParagraphStyle(name="Metric",parent=styles["BodyText"],fontSize=8,leading=11,textColor=colors.HexColor("#173D2B")))
        start_date=self.start_date.get(); end_date=self.end_date.get()
        by_user_day={}; all_by_day={}
        for uid,stamp,state in rows:
            uid=str(uid); day=stamp[:10]
            by_user_day.setdefault((uid,day),[]).append(stamp)
            all_by_day[day]=all_by_day.get(day,0)+1
        first_day=datetime.strptime(start_date,"%Y-%m-%d").date()
        last_day=datetime.strptime(end_date,"%Y-%m-%d").date()
        if last_day<first_day: raise ValueError("End date must not precede start date.")
        report_dates=[]; cursor=first_day
        while cursor<=last_day:
            report_dates.append(cursor.isoformat()); cursor+=timedelta(days=1)
        with sqlite3.connect(DB) as c: known_ids={str(r[0]) for r in c.execute("SELECT user_id FROM staff")}
        punch_ids={uid for uid,_ in by_user_day}
        staff_ids=sorted(known_ids|punch_ids,key=lambda x:(not x.isdigit(),int(x) if x.isdigit() else x.lower()))
        is_collective=(scope=="All staff"); pagesize=A4; story=[]

        def make_stats(items):
            cells=[]
            for label,value in items:
                cells.append(Paragraph(f"<b>{escape(str(label))}</b><br/>{escape(str(value))}",styles["Metric"]))
            while len(cells)%2: cells.append(Paragraph("",styles["Metric"]))
            data=[cells[i:i+2] for i in range(0,len(cells),2)]
            tab=Table(data,colWidths=[89*mm,89*mm],hAlign="LEFT")
            tab.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor("#EDF4EF")),("BOX",(0,0),(-1,-1),.6,colors.HexColor("#C9D8CD")),("INNERGRID",(0,0),(-1,-1),.4,colors.HexColor("#D7E2DA")),("LEFTPADDING",(0,0),(-1,-1),8),("RIGHTPADDING",(0,0),(-1,-1),8),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6),("VALIGN",(0,0),(-1,-1),"MIDDLE")]))
            return tab

        def make_chart(day_counts):
            ordered=sorted(day_counts.items())[-24:]
            chart=Drawing(180*mm,38*mm)
            bc=VerticalBarChart(); bc.x=18; bc.y=12; bc.height=27*mm; bc.width=158*mm
            if ordered:
                bc.data=[tuple(v for _,v in ordered)]
                bc.categoryAxis.categoryNames=[d[5:] for d,_ in ordered]
                maxv=max(v for _,v in ordered) or 1
                bc.valueAxis.valueMin=0; bc.valueAxis.valueMax=maxv*1.15; bc.valueAxis.valueStep=max(1,int(maxv/4))
                bc.bars[0].fillColor=colors.HexColor("#477C57")
            else:
                bc.data=[(0,)]; bc.categoryAxis.categoryNames=["No records"]
                bc.valueAxis.valueMin=0; bc.valueAxis.valueMax=1; bc.valueAxis.valueStep=1
            bc.categoryAxis.labels.fontSize=5.5; bc.categoryAxis.labels.angle=45; bc.categoryAxis.labels.boxAnchor="ne"; bc.valueAxis.labels.fontSize=6
            chart.add(bc); return chart

        def make_detail(uid):
            widths=[22*mm,13*mm,28*mm,28*mm,15*mm,74*mm]
            detail=[[Paragraph(x,styles["CellHead"]) for x in ["Date","Punches","First punch","Last punch","Hours","Attendance status"]]]
            for day in report_dates:
                stamps=sorted(by_user_day.get((uid,day),[]))
                status,first,last,worked=self._attendance_day_status(uid,day,stamps)
                vals=[day,str(len(stamps)),first or "—",last or "—",f"{worked:.2f}",status]
                detail.append([Paragraph(escape(str(v)),styles["Cell"]) for v in vals])
            tab=Table(detail,repeatRows=1,colWidths=widths,hAlign="LEFT")
            tab.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#24563B")),("GRID",(0,0),(-1,-1),.35,colors.HexColor("#CFD9D1")),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F4F7F5")]),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),4),("RIGHTPADDING",(0,0),(-1,-1),4),("TOPPADDING",(0,0),(-1,-1),4),("BOTTOMPADDING",(0,0),(-1,-1),4)]))
            return tab

        def employee_page(uid):
            entries=[(day,by_user_day.get((uid,day),[])) for day in report_dates]
            employee_rows=[r for r in rows if str(r[0])==uid]
            daily_counts={day:len(stamps) for day,stamps in entries if stamps}
            incomplete=sum(1 for day,stamps in entries if len(stamps)==1)
            active_days=sum(1 for day,stamps in entries if stamps)
            first_stamp=min((r[1] for r in employee_rows),default="—")
            last_stamp=max((r[1] for r in employee_rows),default="—")
            story.append(Paragraph(escape(self._name(uid)),styles["EmployeeTitle"]))
            story.append(Paragraph(f"Employee ID: {escape(uid)} &nbsp; | &nbsp; Period: {start_date} to {end_date}",styles["BodyText"]))
            story.append(Spacer(1,3*mm))
            story.append(make_stats([("TOTAL PUNCHES",len(employee_rows)),("ACTIVE DAYS",active_days),("INCOMPLETE DAYS",incomplete),("FIRST PUNCH",first_stamp[:16] if first_stamp!="—" else "—"),("LAST PUNCH",last_stamp[:16] if last_stamp!="—" else "—")]))
            story.extend([Spacer(1,2*mm),Paragraph("Attendance activity",styles["Heading2"]),make_chart(daily_counts),Spacer(1,1*mm),Paragraph("Daily attendance records",styles["Heading2"]),make_detail(uid),Spacer(1,2*mm),Paragraph("Status is calculated using saved schedules, grace periods, minimum hours, working weekdays and holidays.",styles["SmallNote"])])

        if is_collective:
            story.append(Paragraph("SCAFA Workforce — Collective Attendance Report",styles["ReportTitle"]))
            story.append(Paragraph(f"All staff &nbsp; | &nbsp; Period: {start_date} to {end_date} &nbsp; | &nbsp; Generated: {datetime.now().strftime('%d %b %Y, %H:%M')}",styles["BodyText"]))
            story.append(Spacer(1,4*mm))
            story.append(make_stats([("TOTAL PUNCHES",len(rows)),("STAFF",len(staff_ids)),("ACTIVE DAYS",len(all_by_day)),("AVERAGE PUNCHES / DAY",f"{len(rows)/max(1,len(all_by_day)):.1f}")]))
            story.extend([Spacer(1,3*mm),Paragraph("Collective attendance activity",styles["Heading2"]),make_chart(all_by_day)])
            for uid in staff_ids:
                story.append(PageBreak()); employee_page(uid)
        else: employee_page(str(scope))

        def footer(canvas,doc):
            canvas.saveState(); width=pagesize[0]
            canvas.setStrokeColor(colors.HexColor("#C9D8CD")); canvas.line(15*mm,12*mm,width-15*mm,12*mm)
            canvas.setFont("Helvetica",7); canvas.setFillColor(colors.HexColor("#64748b"))
            canvas.drawString(15*mm,7*mm,"SCAFA Workforce • Confidential attendance record")
            canvas.drawRightString(width-15*mm,7*mm,f"Page {doc.page}")
            canvas.restoreState()
        doc=Doc(path,pagesize=pagesize,rightMargin=15*mm,leftMargin=15*mm,topMargin=13*mm,bottomMargin=17*mm,title="SCAFA Workforce Attendance Report")
        doc.build(story,onFirstPage=footer,onLaterPages=footer)

if __name__=="__main__":
    App().mainloop()